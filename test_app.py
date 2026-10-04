"""Streamlit workflow tests; no browser session or participant data required."""
import logging
import unittest
from pathlib import Path
from streamlit.testing.v1 import AppTest
from test_dra import perfect_session
from scoring import CORRECT, TIMING_COMMISSION, TIMING_OMISSION, NA

class AppWorkflowTests(unittest.TestCase):
 def setUp(self):
  logging.disable(logging.WARNING)
 def tearDown(self):
  logging.disable(logging.NOTSET)
 def app(self):
  return AppTest.from_file(str(Path(__file__).with_name('app.py')),default_timeout=30)
 def assert_no_exception(self,app):
  self.assertFalse(app.exception,[e.message for e in app.exception])
 def test_selection_navigation_and_separate_score_controls(self):
  app=self.app().run();self.assert_no_exception(app)
  next(b for b in app.button if b.label=='Use this set').click().run()
  self.assertEqual(app.session_state['planner_selection']['set_id'],'DRA-001')
  fixed=list(app.session_state['planner_selection']['ordered_scenario_ids'])
  self.assertNotIn('set_id',app.session_state['session'])
  self.assertTrue(all(t['scenario_id'] is None for t in app.session_state['trials']))
  app.checkbox(key='code_1_worksheet_occ_check_0').check().run()
  app.checkbox(key='code_1_worksheet_timing_check_1').check().run()
  app.checkbox(key='code_1_worksheet_direction_check_0').check().run()
  app.radio(key='trial_index').set_value(2).run()
  app.radio(key='trial_index').set_value(1).run();self.assert_no_exception(app)
  self.assertEqual(app.session_state['trials'][0]['actions']['worksheet']['timing'],TIMING_COMMISSION)
  self.assertEqual(app.session_state['planner_selection']['ordered_scenario_ids'],fixed)
  self.assertFalse(any('Expected prompts' in x.value for x in app.info))
 def test_completed_session_final_summary_and_dirty_review(self):
  session,trials=perfect_session()
  session.update(participant_id='P1',session_number='1',data_collector='Mary',collector_role='Primary',
                 date='2026-10-03',end_reason='10 trials completed',end_at='8:00')
  app=self.app();app.session_state['session']=session;app.session_state['trials']=trials
  app.run();self.assert_no_exception(app)
  app.session_state['trials'][0]['reviewed']=True
  app.run();self.assert_no_exception(app)
  self.assertEqual(next(m.value for m in app.metric if m.label=='Overall fidelity'),'100.0%')
  app.checkbox(key='code_1_instruction_timing_check_1').check().run()
  app.checkbox(key='code_1_instruction_direction_check_1').check().run()
  self.assertTrue(any('timing conflicts' in e.value for e in app.error))
  self.assertFalse(app.session_state['trials'][0]['reviewed'])
  self.assertTrue(any(m.label=='Provisional fidelity' for m in app.metric))
 def test_learner_fidelity_stays_separate_and_locks_plan(self):
  app=self.app().run()
  next(b for b in app.button if b.label=='Use this set').click().run()
  app.checkbox(key='leader_trial_1_step_1_check_0').check().run()
  self.assertEqual(app.session_state['learner_fidelity']['1']['scores']['step_1'],'Yes')
  self.assertTrue(next(b for b in app.button if b.label=='Randomly select a session').disabled)
  self.assertTrue(all(a['occurrence']=='Missing' for t in app.session_state['trials'] for a in t['actions'].values()))
  self.assertNotIn('set_id',app.session_state['session'])
  self.assert_no_exception(app)
 def test_default_complete_and_optional_status_buttons(self):
  app=self.app().run()
  self.assertEqual(app.session_state['trials'][0]['trial_status'],'Complete')
  self.assertTrue(any(c.key=='code_1_worksheet_occ_check_0' for c in app.checkbox))
  app.button(key='code_1_partial').click().run()
  self.assertEqual(app.session_state['trials'][0]['trial_status'],'Partial')
  app.button(key='code_1_not_observed').click().run()
  self.assertEqual(app.session_state['trials'][0]['trial_status'],'Not reached')
  app.radio(key='trial_index').set_value(2).run()
  app.radio(key='trial_index').set_value(1).run()
  self.assertEqual(app.session_state['trials'][0]['trial_status'],'Not reached')
  app.button(key='code_1_complete').click().run()
  self.assertEqual(app.session_state['trials'][0]['trial_status'],'Complete')
  self.assert_no_exception(app)
 def test_learner_workbook_and_file_identity(self):
  from planner import fidelity_workbook, fidelity_filename
  from dra import select_session, load_catalog
  from io import BytesIO
  from openpyxl import load_workbook
  selection=select_session(load_catalog(),'DRA-001')
  leader={'participant_id':'P12','session_id':'S3','session_leader':'Mary','simulated_learner':'Sam'}
  records={'1':{'scores':{'step_1':'Yes','step_2':'No'},'notes':'Learner error'}}
  book=load_workbook(BytesIO(fidelity_workbook(selection,leader,records)))
  self.assertEqual(fidelity_filename(leader),'P12_S3_Simulated_Learner_Fidelity')
  self.assertEqual(book['Session Summary']['B2'].value,'P12')
  self.assertEqual(book['Session Summary']['B15'].value,0.5)
  self.assertEqual(book['Session Summary']['B15'].number_format,'0.0%')
  self.assertIn('Trial Checklist',book.sheetnames)
 def test_random_selection_does_not_redraw_on_rerun(self):
  app=self.app().run()
  next(b for b in app.button if b.label=='Randomly select a session').click().run()
  selection=dict(app.session_state['planner_selection'])
  app.run()
  self.assertEqual(app.session_state['planner_selection']['set_id'],selection['set_id'])
  self.assertEqual(app.session_state['planner_selection']['selected_at'],selection['selected_at'])

if __name__=='__main__':unittest.main()
