"""Live-session catalog, persistence, and exports. Scoring is in scoring.py."""
from __future__ import annotations
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import json
import random
from typing import Any
import pandas as pd
from scoring import RULES_VERSION, empty_trials, score_trials

ROOT=Path(__file__).parent
SESSION_SCHEMA_VERSION=2

def load_catalog() -> dict[str,Any]:
 catalog=json.loads((ROOT/'dra_catalog_v1.json').read_text())
 ids=[s['set_id'] for s in catalog['sets']]
 if len(ids)!=324 or len(set(ids))!=324: raise ValueError('Catalog needs 324 distinct set IDs.')
 scenarios={s['scenario_id']:s for s in catalog['scenarios']}
 for item in catalog['sets']:
  order=item['ordered_scenario_ids']
  if len(order)!=10 or len(set(order))!=10 or sorted(order)!=item['scenario_ids']: raise ValueError('Invalid catalog membership/order.')
  rows=[scenarios[x] for x in order]
  if sum(x['completed'] for x in rows)!=5 or sum(x['required_prompts'] for x in rows)!=13: raise ValueError('Invalid opportunity balance.')
  for start in (0,5):
   half=rows[start:start+5]
   if sum(x['completed'] for x in half) not in (2,3) or sum(x['required_prompts']==0 for x in half)!=1: raise ValueError('Invalid half balance.')
   if [sum(x['behavior']==b for x in half) for b in ('No tapping or banging','Finger tapping','Table banging')]!=[1,2,2]: raise ValueError('Invalid behavior balance.')
  for a,b,c in zip(rows,rows[1:],rows[2:]):
   if a['completed']==b['completed']==c['completed'] or a['behavior']==b['behavior']==c['behavior']: raise ValueError('Invalid run length.')
 return catalog

def select_session(catalog: dict[str,Any],set_id: str|None=None) -> dict[str,Any]:
 item=random.SystemRandom().choice(catalog['sets']) if set_id is None else next((s for s in catalog['sets'] if s['set_id']==set_id),None)
 if item is None: raise ValueError('Unknown set ID.')
 return {'catalog_version':catalog['catalog_version'],'set_id':item['set_id'],
         'ordered_scenario_ids':list(item['ordered_scenario_ids']),
         'selected_at':datetime.now(timezone.utc).isoformat(),'mode':'in_person'}

def session_plan(selection: dict[str,Any],catalog: dict[str,Any]) -> list[dict[str,Any]]:
 scenarios={s['scenario_id']:s for s in catalog['scenarios']}
 return [{'Trial':i,'Scenario':sid,'Task pattern':scenarios[sid]['task_pattern'],'Behavior':scenarios[sid]['behavior'],
          'Dolphin earned':scenarios[sid]['completed'],'Prompts expected':scenarios[sid]['required_prompts']}
         for i,sid in enumerate(selection['ordered_scenario_ids'],1)]

def load_guide() -> dict[str,list[dict[str,str]]]:
 entries=json.loads((ROOT/'scoring_guide.json').read_text())['entries'];grouped={}
 for entry in entries: grouped.setdefault(entry['section'],[]).append(entry)
 return grouped

def serialize_session(session: dict[str,Any],setup: dict[str,Any],trials: list[dict[str,Any]]) -> bytes:
 return (json.dumps({'schema_version':SESSION_SCHEMA_VERSION,'rules_version':RULES_VERSION,
                    'session':session,'setup':setup,'trials':trials},indent=2,default=str)+'\n').encode()

def restore_session(data: bytes) -> dict[str,Any]:
 if len(data)>5_000_000: raise ValueError('Session file is too large.')
 payload=json.loads(data)
 if payload.get('schema_version')!=SESSION_SCHEMA_VERSION or payload.get('rules_version')!=RULES_VERSION:
  raise ValueError('This file uses another scoring format. Legacy scores cannot be reinterpreted automatically.')
 trials=payload.get('trials')
 if not isinstance(trials,list) or len(trials)!=10 or [t.get('trial') for t in trials]!=list(range(1,11)):
  raise ValueError('Session must contain trial records 1–10.')
 session=payload.get('session',{})
 if session.get('set_id'):
  canonical=select_session(load_catalog(),session['set_id'])
  if session.get('catalog_version')!=canonical['catalog_version'] or session.get('ordered_scenario_ids')!=canonical['ordered_scenario_ids']:
   raise ValueError('Set/order does not match the fixed catalog.')
  if any(t.get('scenario_id')!=sid for t,sid in zip(trials,canonical['ordered_scenario_ids'])): raise ValueError('Trial scenario differs from stored plan.')
 issues=score_trials(trials)['validation_issues']
 if issues: raise ValueError('Invalid session data: '+issues[0])
 return payload

def summary_row(session: dict[str,Any],scores: dict[str,Any],*,final: bool) -> dict[str,Any]:
 return {'Participant ID':session.get('participant_id',''),'Session number':session.get('session_number',''),
         'Date':session.get('date',''),'Collector':session.get('data_collector',''),'Collector role':session.get('collector_role',''),
         'Mode':'in_person','Set ID':session.get('set_id',''),'Catalog version':session.get('catalog_version',''),
         'Rules version':RULES_VERSION,'Trial order':', '.join(map(str,session.get('ordered_scenario_ids',[]))),
         'Status':'Final' if final else 'Draft','Correct scores':scores['correct'],'Applicable scores':scores['applicable'],
         'Overall fidelity (%)':scores['percent'] if final else None,'Provisional fidelity (%)':scores['percent'] if not final else None,
         'Missing scores':scores['missing'],**scores['counts']}


def session_results_rows(trials: list[dict[str,Any]],scores: dict[str,Any]) -> list[dict[str,Any]]:
 """Wide, human-readable coding view; canonical scores remain authoritative."""
 canonical={(r['trial'],r['component'],r['measure']):r['result'] for r in scores['details']}
 by_trial={t['trial']:t for t in trials}
 rows=[]
 def add(label,getter):
  rows.append({'Behavior / step':label,**{f'Trial {i}':getter(by_trial[i]) if i in by_trial else 'Not observed' for i in range(1,11)}})
 def observation(t,key):
  if t.get('trial_status')=='Not reached': return 'Not observed'
  value=t.get('observations',{}).get(key)
  if isinstance(value,bool): return 'Yes' if value else 'No'
  return value if value is not None else 'Missing'
 def response(t,key,measure,inverse=False):
  if t.get('trial_status')=='Not reached': return 'Not observed'
  result=canonical.get((t['trial'],key,measure),'Missing')
  if result=='Correct': return 'No' if inverse else 'Yes'
  if result in ('Omission','Commission'): return 'Yes' if inverse else 'No'
  if result=='Timing commission': return 'No (early)'
  if result=='Timing omission': return 'No (late)'
  if result=='Terminated by participant': return 'Ended early by participant'
  return result
 add('Trial status',lambda t:'Not observed' if t.get('trial_status')=='Not reached' else t.get('trial_status'))
 add('Problems completed',lambda t:'Not observed' if t.get('trial_status')=='Not reached' else ('N/A' if t.get('observations',{}).get('problems_na') else observation(t,'problems')))
 add('Finger tapping occurred',lambda t:observation(t,'tapping_observed'))
 add('Table banging occurred',lambda t:observation(t,'banging_observed'))
 items=[
 ('worksheet','Worksheet presented','Worksheet presented within 3 seconds'),
 ('instruction','Initial instruction given','Initial instruction given within 3 seconds'),
 ('prompt_1','First required prompt given','First prompt given on time'),
 ('prompt_2','Second required prompt given','Second prompt given on time'),
 ('worksheet_removal','Worksheet removed','Worksheet removed on time'),
 ('earned_dolphin','Dolphin given after two problems completed','Earned dolphin given within 3 seconds'),
 ('dolphin_removal','Dolphin removed after earned access','Earned dolphin access lasted 13–17 seconds')]
 for key,label,timing_label in items:
  add(label,lambda t,k=key:response(t,k,'Occurrence'))
  add(timing_label,lambda t,k=key:response(t,k,'Timing'))
 for key,label,inverse in [
 ('no_unearned','Dolphin given before two problems completed',True),
 ('no_excess','No more than two prompts given',False),
 ('no_work_prompts','No prompts given during ongoing work',False),
 ('no_tapping_comments','No stop statements for finger tapping',False),
 ('no_banging_comments','No stop statements for table banging',False)]:
  add(label,lambda t,k=key,inv=inverse:response(t,k,'Behavior',inv))
 add('Timer used (descriptive)',lambda t:observation(t,'timer_use'))
 add('Trial notes',lambda t:t.get('notes',''))
 return rows

def make_workbook(session: dict[str,Any],setup: dict[str,Any],trials: list[dict[str,Any]],scores: dict[str,Any],*,final: bool=False) -> bytes:
 """Runtime export preserves the app's pandas/openpyxl implementation."""
 output=BytesIO();observations=[];actions=[];events=[]
 canonical={(r['trial'],r['component'],r['measure']):r['result'] for r in scores['details']}
 for trial in trials:
  observations.append({'Trial':trial['trial'],'Scenario':trial.get('scenario_id'),'Status':trial.get('trial_status'),**trial.get('observations',{}),'Notes':trial.get('notes','')})
  actions.extend({'Trial':trial['trial'],'Scenario':trial.get('scenario_id'),'Component':key,**action,
                  'occurrence':canonical.get((trial['trial'],key,'Occurrence'),action.get('occurrence')),
                  'timing':canonical.get((trial['trial'],key,'Timing'),action.get('timing'))}
                 for key,action in trial.get('actions',{}).items())
  events.extend({'Trial':trial['trial'],'Scenario':trial.get('scenario_id'),**event} for event in trial.get('events',[]))
 guide=[r for entries in load_guide().values() for r in entries]
 with pd.ExcelWriter(output,engine='openpyxl') as writer:
  summary=summary_row(session,scores,final=final)
  for field in ('Trial order','Status','Provisional fidelity (%)'):
   summary.pop(field,None)
  denominator=scores['applicable']
  counts=scores['counts']
  summary['Commission error percentage']=100*counts.get('Commission',0)/denominator if final and denominator else None
  summary['Omission error percentage']=100*counts.get('Omission',0)/denominator if final and denominator else None
  summary['Timing error percentage']=100*(counts.get('Timing commission',0)+counts.get('Timing omission',0))/denominator if final and denominator else None
  pd.DataFrame([summary]).to_excel(writer,sheet_name='Session Summary',index=False)
  pd.DataFrame(session_results_rows(trials,scores)).to_excel(writer,sheet_name='Session Results',index=False)
  pd.DataFrame([{'Field':k,'Value':json.dumps(v,default=str) if isinstance(v,(list,dict)) else str(v)} for k,v in session.items()]).to_excel(writer,sheet_name='Session',index=False)
  pd.DataFrame([{'Field':k,'Value':v} for k,v in setup.items()]).to_excel(writer,sheet_name='Setup Observations',index=False)
  if session.get('set_id'): pd.DataFrame(session_plan(session,load_catalog())).to_excel(writer,sheet_name='Fixed Trial Plan',index=False)
  pd.DataFrame(observations).to_excel(writer,sheet_name='Observations',index=False)
  pd.DataFrame(actions).to_excel(writer,sheet_name='Action Records',index=False)
  pd.DataFrame(scores['step_summary']).to_excel(writer,sheet_name='Step Summaries',index=False)
  pd.DataFrame(scores['details']).to_excel(writer,sheet_name='Scoring Details',index=False)
  pd.DataFrame(events,columns=['Trial','Scenario','Time','Event','Classification','Notes']).to_excel(writer,sheet_name='Event Log',index=False)
  pd.DataFrame(guide).to_excel(writer,sheet_name='Coding Instructions',index=False)
  for sheet in writer.book.worksheets:
   sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
   for column in sheet.columns:
    sheet.column_dimensions[column[0].column_letter].width=max(14,min(max(len(str(c.value or '')) for c in column)+2,70))
    for cell in column:
     from openpyxl.styles import Alignment
     cell.alignment=Alignment(vertical='top',wrap_text=True)
     if isinstance(cell.value,str) and cell.value.startswith('='): cell.data_type='s'

  sheet=writer.book['Session Results']
  sheet.freeze_panes='B2';sheet.auto_filter.ref=None
  sheet.column_dimensions['A'].width=54
  from openpyxl.styles import Font, PatternFill
  for cell in sheet[1]:
   cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='243F54')
  for row in sheet.iter_rows(min_row=2):
   sheet.row_dimensions[row[0].row].height=36
   row[0].font=Font(bold=True)
   for cell in row[1:]:
    sheet.column_dimensions[cell.column_letter].width=18
    cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
   if row[0].row%2==0:
    for cell in row: cell.fill=PatternFill('solid',fgColor='EDF2F5')
 return output.getvalue()
